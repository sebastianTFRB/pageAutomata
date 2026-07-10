import csv
import os
import sqlite3
import unicodedata
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from modules.logger import log


class ExcelManager:

    DEVICE_TYPE_ALIASES = {
        "all": "all",
        "todos": "all",
        "switch": "switch",
        "switches": "switch",
        "sw": "switch",
        "camara": "camara",
        "camaras": "camara",
        "camera": "camara",
        "cameras": "camara",
        "workstation": "workstation",
        "workstations": "workstation",
        "videowall": "videowall",
        "video wall": "videowall",
        "rack": "rack",
        "racks": "rack",
    }

    def __init__(self, archivo):

        self.archivo = str(archivo)
        self.wb = None
        self.ws = None
        self.conn = None
        self.cursor = None
        self.is_sqlite = self.archivo.lower().endswith((".db", ".sqlite", ".sqlite3"))

    def __del__(self):

        self.cerrar()

    def _crear_directorio(self):

        directorio = os.path.dirname(self.archivo)

        if directorio:
            os.makedirs(directorio, exist_ok=True)

    def _crear_tabla(self):

        self.cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS switches (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                tipo TEXT,
                nombre TEXT,
                ip TEXT,
                usuario TEXT,
                password TEXT,
                zona TEXT,
                estado TEXT,
                ssh TEXT,
                spanning_tree TEXT,
                fecha TEXT,
                falla TEXT,
                observacion TEXT
            )
            """
        )

        # Migra columnas faltantes en bases antiguas.
        self._asegurar_columna("switches", "tipo", "TEXT")
        self._asegurar_columna("switches", "zona", "TEXT")
        self._asegurar_columna("switches", "spanning_tree", "TEXT")
        self._asegurar_columna("switches", "falla", "TEXT")
        self._asegurar_columna("switches", "observacion", "TEXT")

        self.cursor.execute(
            """
            CREATE TABLE IF NOT EXISTS ssh_command_library (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                comando TEXT NOT NULL UNIQUE,
                descripcion TEXT,
                created_at TEXT NOT NULL
            )
            """
        )

        self.conn.commit()

    def _asegurar_columna(self, tabla, columna, tipo_sql):

        self.cursor.execute(f"PRAGMA table_info({tabla})")
        columnas = [fila[1] for fila in self.cursor.fetchall()]

        if columna in columnas:
            return

        self.cursor.execute(f"ALTER TABLE {tabla} ADD COLUMN {columna} {tipo_sql}")

    def abrir(self):

        self._crear_directorio()

        if self.is_sqlite:

            log.info(f"Abriendo SQLite: {self.archivo}")

            self.conn = sqlite3.connect(self.archivo)
            self.conn.row_factory = sqlite3.Row
            self.cursor = self.conn.cursor()
            self._crear_tabla()

            return

        log.info(f"Abriendo Excel: {self.archivo}")

        if os.path.exists(self.archivo):
            self.wb = load_workbook(self.archivo)
        else:
            self.wb = Workbook()
            self.ws = self.wb.active
            self.ws.title = "Switches"
            self.ws.append([
                "id",
                "nombre",
                "ip",
                "usuario",
                "password",
                "zona",
                "estado",
                "ssh",
                "spanning_tree",
                "fecha",
                "falla",
                "observacion"
            ])
            self.wb.save(self.archivo)

        self.ws = self.wb.active

    def _normalizar(self, valor):

        if valor is None:
            return ""

        return str(valor).strip()

    def _normalizar_tipo(self, tipo, default="switch"):

        valor = self._normalizar(tipo).lower()

        if not valor:
            valor = default

        valor = unicodedata.normalize("NFKD", valor).encode("ascii", "ignore").decode("ascii")

        return self.DEVICE_TYPE_ALIASES.get(valor, valor or default)

    def _insertar_dispositivo(self, datos):

        tipo = self._normalizar_tipo(datos.get("tipo"), default="switch")

        self.cursor.execute(
            """
            INSERT INTO switches (
                tipo,
                nombre,
                ip,
                usuario,
                password,
                zona,
                estado,
                ssh,
                spanning_tree,
                fecha,
                falla,
                observacion
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                tipo,
                self._normalizar(datos.get("nombre")),
                self._normalizar(datos.get("ip")),
                self._normalizar(datos.get("usuario")),
                self._normalizar(datos.get("password")),
                self._normalizar(datos.get("zona")),
                self._normalizar(datos.get("estado") or "en funcionamiento"),
                self._normalizar(datos.get("ssh")),
                self._normalizar(datos.get("spanning_tree")),
                self._normalizar(datos.get("fecha") or datetime.now().date().isoformat()),
                self._normalizar(datos.get("falla")),
                self._normalizar(datos.get("observacion"))
            )
        )

    def importar_desde_csv(self, ruta_csv, tipo_default="switch"):

        if self.conn is None:
            self.abrir()

        with open(ruta_csv, newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle)

            for row in reader:
                if not row.get("tipo"):
                    row["tipo"] = tipo_default
                self._insertar_dispositivo(row)

        self.conn.commit()
        self.cerrar()

        log.info(f"CSV importado desde {ruta_csv}")

    def importar_desde_excel(self, ruta_excel, tipo_default="switch"):

        if self.conn is None:
            self.abrir()

        workbook = load_workbook(ruta_excel)
        sheet = workbook.active

        headers = {}
        for col_idx, cell in enumerate(sheet[1], 1):
            if cell.value:
                headers[str(cell.value).lower().strip()] = col_idx

        log.info(f"Headers detectados: {headers}")
        log.info(f"Filas totales en Excel: {sheet.max_row}")

        contador = 0
        for fila in range(2, sheet.max_row + 1):
            datos = {}

            for campo in ["tipo", "nombre", "ip", "usuario", "password", "zona", "estado", "ssh", "spanning_tree", "fecha", "falla", "observacion"]:
                col_idx = headers.get(campo)
                if col_idx is None:
                    for key in headers:
                        if campo in key.lower():
                            col_idx = headers[key]
                            break
                
                if col_idx:
                    valor = sheet.cell(row=fila, column=col_idx).value
                    datos[campo] = valor if valor else ""
                else:
                    datos[campo] = ""

            if not datos.get("tipo"):
                datos["tipo"] = tipo_default

            ip_valor = datos.get("ip")
            log.info(f"Fila {fila}: IP={ip_valor}, Nombre={datos.get('nombre')}")

            if not ip_valor:
                continue

            self._insertar_dispositivo(datos)
            contador += 1

        self.conn.commit()
        self.cerrar()

        log.info(f"Excel importado desde {ruta_excel} - {contador} registros insertados")

    def importar_desde_archivo(self, ruta_archivo, tipo_default="switch"):

        extension = Path(ruta_archivo).suffix.lower()

        if extension == ".csv":
            return self.importar_desde_csv(ruta_archivo, tipo_default=tipo_default)

        return self.importar_desde_excel(ruta_archivo, tipo_default=tipo_default)

    def obtener_dispositivos(self, tipo=None):

        if self.is_sqlite:
            if self.conn is None:
                self.abrir()

            tipo_normalizado = self._normalizar_tipo(tipo, default="all")

            if tipo_normalizado == "all":
                self.cursor.execute(
                    """
                    SELECT id, COALESCE(NULLIF(tipo, ''), 'switch') AS tipo, nombre, ip, usuario, password, zona, estado, ssh, spanning_tree, fecha, falla, observacion
                    FROM switches
                    ORDER BY id
                    """
                )
            else:
                self.cursor.execute(
                    """
                    SELECT id, COALESCE(NULLIF(tipo, ''), 'switch') AS tipo, nombre, ip, usuario, password, zona, estado, ssh, spanning_tree, fecha, falla, observacion
                    FROM switches
                    WHERE LOWER(COALESCE(NULLIF(tipo, ''), 'switch')) = ?
                    ORDER BY id
                    """,
                    (tipo_normalizado,)
                )

            filas = self.cursor.fetchall()

            dispositivos = []

            for fila in filas:
                dispositivos.append({
                    "id": fila[0],
                    "tipo": self._normalizar(fila[1]) or "switch",
                    "nombre": self._normalizar(fila[2]),
                    "ip": self._normalizar(fila[3]),
                    "usuario": self._normalizar(fila[4]),
                    "password": self._normalizar(fila[5]),
                    "zona": self._normalizar(fila[6]),
                    "estado": self._normalizar(fila[7]),
                    "ssh": self._normalizar(fila[8]),
                    "spanning_tree": self._normalizar(fila[9]),
                    "fecha": self._normalizar(fila[10]),
                    "falla": self._normalizar(fila[11]),
                    "observacion": self._normalizar(fila[12])
                })

            log.info(f"Se encontraron {len(dispositivos)} dispositivos")

            self.cerrar()

            return dispositivos

        if self.ws is None:
            self.abrir()

        dispositivos = []

        for fila in range(2, self.ws.max_row + 1):

            ip = self.ws[fila][4].value          # Columna E
            usuario = self.ws[fila][8].value    # Columna I
            password = self.ws[fila][9].value   # Columna J
            nombre = self.ws[fila][12].value    # Columna M

            if not ip:
                continue

            dispositivos.append({

                "fila": fila,
                "tipo": "switch",

                "nombre": nombre,

                "ip": str(ip).strip(),

                "usuario": str(usuario).strip(),

                "password": str(password).strip(),
                "zona": "",
                "estado": "en funcionamiento",
                "ssh": "",
                "spanning_tree": "",
                "fecha": ""

            })

        log.info(f"Se encontraron {len(dispositivos)} dispositivos")

        return dispositivos

    def obtener_switches(self):

        return self.obtener_dispositivos("switch")

    def actualizar_resultado(

        self,
        fila,
        estado,
        falla="",
        observacion=""
    ):

        if self.is_sqlite:
            if self.conn is None:
                self.abrir()

            self.cursor.execute(
                "UPDATE switches SET estado = ?, falla = ?, observacion = ? WHERE id = ?",
                (estado, falla, observacion, fila)
            )
            self.conn.commit()
            return

        self.ws[fila][13].value = estado        # N
        self.ws[fila][14].value = falla         # O
        self.ws[fila][15].value = observacion   # P

    def exportar_csv(self, ruta_csv):

        if self.conn is None:
            self.abrir()

        dispositivos = self.obtener_switches()

        with open(ruta_csv, "w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=["id", "tipo", "nombre", "ip", "usuario", "password", "estado", "ssh", "falla", "observacion", "fecha"]
            )
            writer.writeheader()

            for switch in dispositivos:
                writer.writerow({
                    "id": switch.get("id", ""),
                    "tipo": switch.get("tipo", "switch"),
                    "nombre": switch.get("nombre", ""),
                    "ip": switch.get("ip", ""),
                    "usuario": switch.get("usuario", ""),
                    "password": switch.get("password", ""),
                    "estado": switch.get("estado", ""),
                    "ssh": switch.get("ssh", ""),
                    "falla": switch.get("falla", ""),
                    "observacion": switch.get("observacion", ""),
                    "fecha": switch.get("fecha", "")
                })

        self.cerrar()

        log.info(f"CSV exportado hacia {ruta_csv}")

    def exportar_excel(self, ruta_excel):

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Dispositivos"
        sheet.append(["id", "tipo", "nombre", "ip", "usuario", "password", "estado", "ssh", "falla", "observacion", "fecha"])

        for switch in self.obtener_switches():
            sheet.append([
                switch.get("id", ""),
                switch.get("tipo", "switch"),
                switch.get("nombre", ""),
                switch.get("ip", ""),
                switch.get("usuario", ""),
                switch.get("password", ""),
                switch.get("estado", ""),
                switch.get("ssh", ""),
                switch.get("falla", ""),
                switch.get("observacion", ""),
                switch.get("fecha", "")
            ])

        workbook.save(ruta_excel)

        self.cerrar()

        log.info(f"Excel exportado hacia {ruta_excel}")

    def cerrar(self):

        if self.conn is not None:
            try:
                self.conn.commit()
            except sqlite3.Error:
                pass

            try:
                self.conn.close()
            except sqlite3.Error:
                pass

            self.conn = None
            self.cursor = None

    def agregar_dispositivo(self, tipo, nombre, ip, usuario="", password="", zona="", estado="en funcionamiento", ssh="", spanning_tree="", fecha="", falla="", observacion=""):

        if self.conn is None:
            self.abrir()

        datos = {
            "tipo": tipo,
            "nombre": nombre,
            "ip": ip,
            "usuario": usuario,
            "password": password,
            "zona": zona,
            "estado": estado,
            "ssh": ssh,
            "spanning_tree": spanning_tree,
            "fecha": fecha or datetime.now().date().isoformat(),
            "falla": falla,
            "observacion": observacion
        }

        self._insertar_dispositivo(datos)
        self.conn.commit()

        log.info(f"Dispositivo agregado: {nombre} ({ip})")

    def agregar_switch(self, nombre, ip, usuario="", password="", zona="", estado="en funcionamiento", ssh="", spanning_tree="", fecha=""):

        self.agregar_dispositivo(
            tipo="switch",
            nombre=nombre,
            ip=ip,
            usuario=usuario,
            password=password,
            zona=zona,
            estado=estado,
            ssh=ssh,
            spanning_tree=spanning_tree,
            fecha=fecha,
        )

    def obtener_comandos_ssh(self):

        if not self.is_sqlite:
            return []

        if self.conn is None:
            self.abrir()

        self.cursor.execute(
            """
            SELECT id, comando, descripcion, created_at
            FROM ssh_command_library
            ORDER BY comando COLLATE NOCASE
            """
        )

        filas = self.cursor.fetchall()

        comandos = []
        for fila in filas:
            comandos.append(
                {
                    "id": fila[0],
                    "comando": self._normalizar(fila[1]),
                    "descripcion": self._normalizar(fila[2]),
                    "created_at": self._normalizar(fila[3]),
                }
            )

        return comandos

    def agregar_comando_ssh(self, comando, descripcion=""):

        if not self.is_sqlite:
            return

        comando_limpio = self._normalizar(comando)
        descripcion_limpia = self._normalizar(descripcion)

        if not comando_limpio:
            raise ValueError("El comando no puede estar vacio")

        if self.conn is None:
            self.abrir()

        self.cursor.execute(
            """
            INSERT OR IGNORE INTO ssh_command_library (comando, descripcion, created_at)
            VALUES (?, ?, ?)
            """,
            (
                comando_limpio,
                descripcion_limpia,
                datetime.now().isoformat(timespec="seconds"),
            ),
        )
        self.conn.commit()

        if self.cursor.rowcount == 0:
            raise ValueError("Ese comando ya existe en la libreria")

        log.info(f"Comando agregado a libreria SSH: {comando_limpio}")

    def guardar(self):

        if self.is_sqlite:
            if self.conn is not None:
                self.conn.commit()

            log.info("SQLite guardado")
            return

        self.wb.save(self.archivo)

        log.info("Excel guardado")